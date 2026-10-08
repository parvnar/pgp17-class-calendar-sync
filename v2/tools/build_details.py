"""Build the Session Details workbook (Sessions + Materials tabs) from the prep plan."""
import re, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
ws = wb['All Sessions']

SUBJ = {'DA': 'DA', 'FLP': 'FLP', 'IK': 'IK', 'MANAC': 'MANAC', 'MComm': 'MCOMM',
        'MM-II': 'MM II', 'ODD': 'ODD', 'QM-II': 'QM II', 'WIPS': 'WIPS', 'MEB': 'MEB'}

def subj_code(name):
    m = re.search(r'\(([^)]+)\)\s*$', name)
    return SUBJ[m.group(1)]

def sessions_of(s):
    s = str(s).strip()
    if '-' in s:
        a, b = s.split('-')
        return list(range(int(a), int(b) + 1))
    return [int(s)]

def clean(v):
    return '' if v is None else str(v).strip()

# Manual clean-ups of placeholder text so every material carries a code.
FIX = {
    ('MEB', 'C', '[Per TOC: C3 Eepnagar and R1 run to session 4]'):
        'C3: The National Income Accounts of Eepnagar (continued – per TOC; confirm with faculty)',
    ('MEB', 'R', '[R1 per TOC]'): 'R1 (continued – per TOC)',
    ('MEB', 'R', '[R6: Much Ado About Multipliers, The Economist, 24 Sep 2009 – per TOC]'):
        'R6: Much Ado About Multipliers, The Economist (24 Sep 2009) – listed here per TOC',
    ('DA', 'C', 'Spreadsheet modelling to determine the optimum hotel room rate for a short high-demand period (INFORMS Transactions on Education, Vol 11, No 1, pp 35-42)'):
        'X1: Spreadsheet modelling to determine the optimum hotel room rate for a short high-demand period (INFORMS Transactions on Education, Vol 11, No 1, pp 35-42)',
    ('ODD', 'C', 'Simulation: Judgment in a Crisis (HBR)'): 'SIM1: Judgment in a Crisis – simulation (HBR)',
}
tn_counter = {}

def split_items(subject, kind, text):
    """Split a case/reading cell into items; returns {session or None: [items]}."""
    text = FIX.get((subject, kind, text), text)
    out = {}
    for line in text.split('\n'):
        line = line.strip()
        if not line:
            continue
        sess = None
        m = re.match(r'^Session (\d+):\s*(.*)$', line)
        if m:
            sess, line = int(m.group(1)), m.group(2)
        m = re.search(r'\s*\(session (\d+)\)\s*$', line)
        if m:
            sess, line = int(m.group(1)), line[:m.start()]
        parts = re.split(r';\s*(?=(?:[CR]\d+:|Unemployment|TSP))', line)
        for p in parts:
            p = p.strip()
            m = re.match(r'^Technical Note:\s*(.+)$', p)
            if m:
                tn_counter[subject] = tn_counter.get(subject, 0) + 1
                p = 'TN%d: %s (Technical Note)' % (tn_counter[subject], m.group(1))
            out.setdefault(sess, []).append(p)
    return out

rows = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    subject = subj_code(r[0])
    sess_list = sessions_of(r[1])
    topic, book, case, reading, notes = (clean(x) for x in r[2:7])
    if subject == 'MEB' and sess_list == [4]:
        topic = 'National Income Accounting (continued). Session 4 is missing from the outline\'s session plan – the TOC runs C3 (Eepnagar) and R1 to session 4.'
    case_items = split_items(subject, 'C', case) if case else {}
    read_items = split_items(subject, 'R', reading) if reading else {}
    for s in sess_list:
        def pick(items):
            if not items:
                return ''
            if any(k is not None for k in items):
                return '\n'.join(items.get(s, []))
            return '\n'.join(items[None])
        rows.append([subject, s, '', topic, book, pick(case_items), pick(read_items), notes])

# Materials: one row per coded item.
materials = {}
CODE = re.compile(r'^\[?((?:C|R|TN|SIM|X)\d+)\b\s*:?\s*(.*)$')
for subject, s, _, _, _, case, reading, _ in rows:
    for kind, cell in (('Case', case), ('Reading', reading)):
        for item in filter(None, cell.split('\n')):
            m = CODE.match(item)
            if not m:
                continue
            code, title = m.group(1), m.group(2).strip()
            key = (subject, code)
            if title.startswith('(contin') or not title:
                title = ''
            if key not in materials or (not materials[key][1] and title):
                materials[key] = [kind, title, materials.get(key, [None, None, set()])[2]]
            materials[key][2].add(s)

out = openpyxl.Workbook()
HDR = Font(bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1F4E78')
WRAP = Alignment(wrap_text=True, vertical='top')

def header(sheet, cols, widths):
    sheet.append(cols)
    for i, c in enumerate(sheet[1]):
        c.font, c.fill = HDR, FILL
        sheet.column_dimensions[c.column_letter].width = widths[i]
    sheet.freeze_panes = 'A2'

sh = out.active
sh.title = 'Sessions'
header(sh, ['Subject', 'Session', 'Faculty (optional)', 'Topic', 'Book & Chapter', 'Case', 'Reading', 'Notes'],
       [10, 9, 12, 60, 28, 55, 55, 40])
for row in rows:
    sh.append(row)
for row in sh.iter_rows(min_row=2):
    for c in row:
        c.alignment = WRAP

mt = out.create_sheet('Materials')
header(mt, ['Subject', 'Code', 'Type', 'Title', 'Sessions', 'Link', 'Link Source', 'Matched File'],
       [10, 8, 9, 70, 12, 45, 12, 45])
order = {'C': 0, 'R': 1, 'TN': 2, 'SIM': 3, 'X': 4}
def mkey(k):
    subject, code = k
    p = re.match(r'([A-Z]+)(\d+)', code)
    return (subject, order[p.group(1)], int(p.group(2)))
for key in sorted(materials, key=mkey):
    kind, title, sess = materials[key]
    mt.append([key[0], key[1], kind, title, ', '.join(str(x) for x in sorted(sess)), '', '', ''])
for row in mt.iter_rows(min_row=2):
    for c in row:
        c.alignment = WRAP

# Keep the original flags for reference.
flags = out.create_sheet('Notes & Flags')
for r in wb['Notes & Flags'].iter_rows(values_only=True):
    flags.append(list(r))
flags.column_dimensions['A'].width = 40
flags.column_dimensions['B'].width = 120
for row in flags.iter_rows():
    for c in row:
        c.alignment = WRAP
for c in flags[1]:
    c.font, c.fill = HDR, FILL

help_ = out.create_sheet('How to edit', 0)
help_.column_dimensions['A'].width = 120
lines = [
    'PGP17 Term II – Session Details (read by the calendar sync script every hour)',
    '',
    'Sessions tab: one row per session. Subject codes match the timetable (MANAC, DA, MEB, QM II, MM II, MCOMM, WIPS, IK, ODD, FLP).',
    '  • Faculty (optional): leave blank for normal rows. Fill a faculty code (e.g. DS) only to give one faculty a different plan for that session – that row then wins for their sections.',
    '  • Case / Reading: one item per line. Start each item with its code (C3: …, R1: …) so the script can attach the Drive link from the Materials tab.',
    '',
    'Materials tab: one row per case / reading / technical note.',
    '  • Link: the script fills this automatically from the Term 2 Drive folder (Link Source = auto). ',
    '  • To fix or add a link yourself, paste it into Link and leave Link Source blank or type "manual" – the script never overwrites manual links.',
    '  • To make the script re-match a row, clear both Link and Link Source.',
    '',
    'Calendar Links tab: created by the script – the add-to-calendar link for each section and the time of the last sync.',
    'Drive Index tab: created by the script – every file found in the Drive folder and which material it was matched to.',
]
for l in lines:
    help_.append([l])
help_['A1'].font = Font(bold=True, size=13)

out.save(OUT)
print('sessions', len(rows), 'materials', len(materials))
